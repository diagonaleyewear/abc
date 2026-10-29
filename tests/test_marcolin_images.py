from datetime import date
import random

from io import BytesIO

import pytest
from openpyxl import load_workbook

from marcolin_images import DEFAULTS, MarcolinConnector, MarcolinPortal, map_portal_status
from provider_images import IMPORTED, NEEDS_REVIEW, PENDING, PortalPackage, ProviderService, ProviderStore
from image_repository import ImageRepository
from provider_support import EANS, fake_odoo, pfile, zip_bytes, photo


class Driver:
    def __init__(self, statuses, package):
        self.statuses, self.package, self.sent = list(statuses), package, []

    def submit(self, eans):
        self.sent.append(list(eans))
        return f"M-{len(self.sent)}"

    def status_text(self, remote_id):
        return self.statuses.pop(0) if len(self.statuses) > 1 else self.statuses[0]

    def download(self, remote_id):
        return self.package


@pytest.fixture
def setup(tmp_path, monkeypatch):
    monkeypatch.setenv("MARCOLIN_AUTOMATION_ENABLED", "true")
    clock = [1_800_000_000.0]
    store = ProviderStore(tmp_path / "s", clock=lambda: clock[0])
    repo = ImageRepository(tmp_path / "i", state=tmp_path / "r", backups=tmp_path / "b")
    from dataclasses import replace
    settings = replace(DEFAULTS, supplier_ids=(573,), max_eans_per_request=2, poll_base_seconds=60)
    service = ProviderService(store, repo, settings, rng=random.Random(0))
    odoo = fake_odoo({573: [(1, "P1", EANS[0]), (2, "P2", EANS[1]), (3, "P3", EANS[2])]})
    return MarcolinConnector(service, odoo), store, repo, clock


def test_status_wording_mapping():
    assert map_portal_status("Pendiente").state == "pending"
    assert map_portal_status(" Disponible ").state == "available"
    assert map_portal_status("Expired").state == "failed"
    assert map_portal_status("algo nuevo").state == "unknown"


def test_deferred_requests_are_tracked_without_duplicates(setup):
    connector, store, repo, clock = setup
    package = PortalPackage(files=[pfile(f"{EANS[0]}.png", view="frontal"), pfile(f"{EANS[1]}.png", view="frontal")])
    driver = Driver(["Pendiente", "Pendiente", "Disponible"], package)
    portal = MarcolinPortal(driver)
    first = connector.detect(date(2026, 9, 1), date(2026, 9, 30))
    assert len(first["created"]) == 2
    connector.service.run_cycle(portal, dry_run=False)
    assert len(driver.sent) == 2 and {r["state"] for r in store.list_requests("marcolin")} == {PENDING}
    # Days later the same orders are detected again: nothing new is sent.
    clock[0] += 3 * 86400
    assert connector.detect(date(2026, 9, 1), date(2026, 9, 30))["created"] == []
    connector.service.run_cycle(portal, dry_run=False)
    clock[0] += 6 * 3600
    connector.service.run_cycle(portal, dry_run=False)
    states = {r["id"]: r["state"] for r in store.list_requests("marcolin")}
    assert IMPORTED in states.values()
    assert {r.ean for r in repo.records()} == {EANS[0], EANS[1]}


def test_unknown_portal_status_requires_review(setup):
    connector, store, _, clock = setup
    connector.detect(date(2026, 9, 1), date(2026, 9, 30), limit=1)
    portal = MarcolinPortal(Driver(["???"], PortalPackage()))
    connector.service.run_cycle(portal, dry_run=False)
    clock[0] += 86400
    connector.service.run_cycle(portal, dry_run=False)
    assert store.list_requests("marcolin")[0]["state"] == NEEDS_REVIEW


def test_manual_assisted_path_without_portal(setup):
    connector, store, repo, _ = setup
    created = connector.detect(date(2026, 9, 1), date(2026, 9, 30), mode="manual")["created"]
    assert connector.service.run_cycle(MarcolinPortal(Driver(["x"], PortalPackage())), dry_run=False)["submitted"] == []
    rid = created[0]
    archive = zip_bytes({"Assets/a.png": photo(), "Assets/manifest.csv":
                         f"filename,ean,view\na.png,{EANS[0]},perspectiva\n".encode()})
    assert connector.import_zip(rid, archive)["imported"] == 1
    assert repo.records(EANS[0])[0].view == "perspectiva"


def test_marcolin_exports_one_code_per_line_and_others_keep_commas(setup):
    connector, store, repo, clock = setup
    upcs = [e[1:] for e in EANS if e.startswith("0")] or [EANS[0]]
    result = connector.plan_eans(upcs[:2] if len(upcs) > 1 else [EANS[0], EANS[1]])
    request_id = result["created"][0]
    exported = connector.export_ean_list(request_id)
    assert "," not in exported and " " not in exported
    assert exported.split("\n") == [e["ean"] for e in store.get(request_id)["eans"]]
    assert connector.export_ean_list(request_id, ",").count(",") == exported.count("\n")
    from provider_images import ProviderConnector
    assert ProviderConnector.list_separator == ","


def test_portal_groups_keep_every_upc_and_inactive_references(setup):
    connector, store, repo, clock = setup
    result = connector.plan_eans([EANS[0], EANS[1]])
    request_id = result["created"][0]
    before = store.get(request_id)["state"]
    report = connector.record_portal_groups(request_id, [
        {"portal_ref": "REF-A", "status": "inactive", "upcs": [EANS[0]]},
        {"portal_ref": "REF-B", "status": "inactive", "upcs": [EANS[1], EANS[2]]},
    ])
    assert report["stored"] == 2 and report["not_in_request"] == [EANS[2]] and report["unassigned"] == []
    groups = connector.portal_groups(request_id)
    assert {g["portal_ref"]: g["upcs"] for g in groups} == {"REF-A": [EANS[0]], "REF-B": [EANS[1]]}
    assert all(g["status"] == "inactive" for g in groups)
    assert store.get(request_id)["state"] == before
    again = connector.record_portal_groups(request_id, [{"portal_ref": "REF-A", "status": "inactive", "upcs": [EANS[0]]}])
    assert again["stored"] == 0
    with pytest.raises(ValueError):
        connector.record_portal_groups(request_id, [{"portal_ref": " ", "upcs": [EANS[0]]}])


def test_portal_groups_match_upc12_to_ean13_and_keep_original(setup):
    connector, store, repo, clock = setup
    from provider_support import make_ean
    ean13 = make_ean("003600029145")
    upc = ean13[1:]
    request_id = connector.plan_eans([upc])["created"][0]
    report = connector.record_portal_groups(request_id, [{"portal_ref": "REF-X", "status": "", "upcs": [ean13]}])
    assert report["stored"] == 1 and report["not_in_request"] == []
    assert connector.portal_groups(request_id)[0]["upcs"] == [upc]


def test_portal_status_change_is_kept_as_history_and_current_view_follows_it(setup):
    connector, store, repo, clock = setup
    request_id = connector.plan_eans([EANS[0]])["created"][0]
    first = [{"portal_ref": "REF-A", "status": "inactive", "upcs": [EANS[0]]}]
    assert connector.record_portal_groups(request_id, first)["stored"] == 1
    assert connector.record_portal_groups(request_id, first)["stored"] == 0
    clock[0] += 3600
    later = [{"portal_ref": "REF-A", "status": "active", "upcs": [EANS[0]]}]
    assert connector.record_portal_groups(request_id, later)["stored"] == 1
    assert [(g["status"], g["upcs"]) for g in connector.portal_groups(request_id)] == [("active", [EANS[0]])]
    clock[0] += 3600
    assert connector.record_portal_groups(request_id, first)["stored"] == 1  # back to inactive is a new observation
    assert connector.portal_groups(request_id)[0]["status"] == "inactive"
    with connector.service.store.connect() as db:
        assert db.execute("SELECT COUNT(*) FROM marcolin_portal_groups").fetchone()[0] == 3
    assert store.get(request_id)["state"] == "ready"


def test_marcolin_excel_keeps_pilot_upcs_as_text_and_confirmed_skus(setup):
    connector, store, _, _ = setup
    upcs = ["664689520060", "664689520084", "664689764686", "664689628810",
            "664689628827", "664689897605"]
    request_ids = connector.plan_eans(upcs)["created"]
    for request_id, request_upcs in zip(request_ids, (upcs[:2], upcs[2:4], upcs[4:])):
        store.save_odoo_items(request_id, [
            {"ean": ean, "sku": f"TF-{index:03d}", "product": "Tom Ford",
             "status": "confirmed"} for index, ean in enumerate(request_upcs, upcs.index(request_upcs[0]) + 1)
        ])
    before = [store.get(request_id) for request_id in request_ids]
    rows = []
    for request_id in request_ids:
        sheet = load_workbook(BytesIO(connector.export_request_xlsx(request_id))).active
        rows.extend(list(sheet.iter_rows(min_row=2)))
        assert sheet["A2"].number_format == "@"
    assert {row[0].value for row in rows} == set(upcs)
    assert {row[1].value for row in rows} == {f"TF-{i:03d}" for i in range(1, 7)}
    assert [store.get(request_id) for request_id in request_ids] == before


def test_marcolin_excel_marks_missing_and_ambiguous_references_for_review(setup):
    connector, store, _, _ = setup
    upcs = [EANS[0], EANS[1], EANS[2]]
    request_ids = connector.plan_eans(upcs)["created"]
    store.save_odoo_items(request_ids[0], [
        {"ean": upcs[0], "status": "pending", "reason": "sin correspondencia confirmada en Odoo"},
        {"ean": upcs[1], "status": "requires_review", "reason": "referencias Odoo ambiguas"},
    ])
    rows = {}
    for request_id in request_ids:
        sheet = load_workbook(BytesIO(connector.export_request_xlsx(request_id))).active
        rows.update({row[0].value: row for row in sheet.iter_rows(min_row=2)})
    assert rows[upcs[0]][1].value is None and rows[upcs[0]][9].value == "pending"
    assert rows[upcs[1]][1].value is None and rows[upcs[1]][9].value == "requires_review"
    assert rows[upcs[2]][10].value == "sin correspondencia Odoo confirmada"


def test_marcolin_excel_generation_is_repeatable_without_mutating_request(setup):
    connector, store, _, _ = setup
    request_id = connector.plan_eans([EANS[0]])["created"][0]
    store.save_odoo_items(request_id, [{"ean": EANS[0], "sku": "SKU-1", "status": "confirmed"}])
    before = store.get(request_id)
    first = connector.export_request_xlsx(request_id)
    second = connector.export_request_xlsx(request_id)
    assert first == second
    assert store.get(request_id) == before
