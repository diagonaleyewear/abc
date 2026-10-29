"""Marcolin connector (#59): deferred asset requests with state tracking.

Documented flow (not yet confirmed with the portal): Multimedia → Assets requests → Product
images + brands → EAN list → validation → confirmations → Profile → Requests → download
(hours or days later). The portal driver is NOT implemented: access method, terms,
MFA/CAPTCHA and the ZIP layout are unconfirmed. Everything else works today:
- planning without duplicating live requests (DB-enforced),
- the manual-assisted path (export EAN list, upload the ZIP the operator downloaded),
- a portal adapter contract the real driver must satisfy.
"""
from __future__ import annotations

from contextlib import closing
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from provider_images import (
    ProviderConnector, ProviderPortal, ProviderSettings, PortalPackage, PortalStatus, UnverifiedProviderPortal,
    ean_key,
)

PROVIDER = "marcolin"
DEFAULTS = ProviderSettings(PROVIDER, max_eans_per_request=50, max_requests_per_day=3, pilot_cap=50)
PENDING_WORDS = {"pending", "pendiente", "requested", "solicitado", "processing", "en proceso", "in progress"}
AVAILABLE_WORDS = {"available", "disponible", "ready", "listo", "completed", "completado", "done"}
FAILED_WORDS = {"rejected", "rechazado", "failed", "error", "expired", "caducado", "cancelled", "cancelado"}


def map_portal_status(raw: str) -> PortalStatus:
    """Portal wording -> state. Unrecognised wording is 'unknown' and forces manual review."""
    word = (raw or "").strip().lower()
    if word in AVAILABLE_WORDS:
        return PortalStatus("available", word)
    if word in FAILED_WORDS:
        return PortalStatus("failed", word)
    if word in PENDING_WORDS:
        return PortalStatus("pending", word)
    return PortalStatus("unknown", word)


class MarcolinDriver(ProviderPortal):
    """Contract for the future real driver (HTTP or browser). ``requests()`` must list the
    portal's own pending/available requests so ambiguous submissions can be reconciled."""

    def requests(self) -> list[dict]:
        raise NotImplementedError


class UnverifiedMarcolinPortal(UnverifiedProviderPortal):
    pass


class MarcolinPortal:
    """Adapts a driver to the core contract, normalising status wording."""

    def __init__(self, driver):
        self.driver = driver

    def submit(self, eans: list[str]) -> str:
        return self.driver.submit(eans)

    def status(self, remote_id: str) -> PortalStatus:
        return map_portal_status(self.driver.status_text(remote_id))

    def download(self, remote_id: str) -> PortalPackage:
        return self.driver.download(remote_id)


class MarcolinConnector(ProviderConnector):
    provider = PROVIDER
    default_mode = "portal"
    defaults = DEFAULTS
    # Marcolin's "Select by list" takes one code per line and warns against commas and spaces.
    list_separator = "\n"

    GROUPS_SCHEMA = """
    CREATE TABLE IF NOT EXISTS marcolin_portal_groups(
      request_id INTEGER NOT NULL REFERENCES requests(id), ean TEXT NOT NULL, reported TEXT NOT NULL,
      portal_ref TEXT NOT NULL, portal_status TEXT NOT NULL DEFAULT '', recorded_at REAL NOT NULL);
    """

    def record_portal_groups(self, request_id: int, groups) -> dict:
        """Keep which portal model/colour reference each requested UPC was grouped under.

        ``groups``: [{"portal_ref", "status", "upcs": [...]}] as read from the portal by an operator. The portal
        status is recorded verbatim and never filters anything (an 'inactive' reference is still kept and not
        discarded); the request state is not touched, so a request pending for days is not resubmitted. UPCs
        that are not part of the request are reported, not stored.
        """
        store = self.service.store
        request = store.get(request_id)
        by_key = {ean_key(e["ean"]): e["ean"] for e in request["eans"]}
        stored, outside, assigned = 0, [], {}
        with closing(store.connect()) as connection:
            connection.executescript(self.GROUPS_SCHEMA)
            connection.execute("BEGIN IMMEDIATE")
            for group in groups:
                ref, status = str(group.get("portal_ref", "")).strip(), str(group.get("status", "")).strip()
                if not ref:
                    raise ValueError("portal_ref_vacio")
                for reported in (str(u).strip() for u in group.get("upcs", [])):
                    original = by_key.get(ean_key(reported)) if reported else None
                    if original is None:
                        outside.append(reported)
                        continue
                    latest = connection.execute(
                        "SELECT portal_status FROM marcolin_portal_groups WHERE request_id=? AND ean=? AND portal_ref=? "
                        "ORDER BY recorded_at DESC, rowid DESC LIMIT 1", (request_id, original, ref)).fetchone()
                    if latest is None or latest[0] != status:  # a changed portal status is new history, not an overwrite
                        connection.execute(
                            "INSERT INTO marcolin_portal_groups(request_id,ean,reported,portal_ref,portal_status,"
                            "recorded_at) VALUES(?,?,?,?,?,?)", (request_id, original, reported, ref, status, store.clock()))
                        stored += 1
                    assigned.setdefault(original, set()).add(ref)
            connection.execute("COMMIT")
        return {"stored": stored, "not_in_request": outside,
                "unassigned": [e for e in by_key.values() if e not in assigned],
                "multi_assigned": {e: sorted(r) for e, r in assigned.items() if len(r) > 1}}

    def portal_groups(self, request_id: int) -> list[dict]:
        """Current groups for a request: every UPC (original literal) under the latest verbatim portal status.

        Earlier statuses stay in the table as history."""
        with closing(self.service.store.connect()) as connection:
            connection.executescript(self.GROUPS_SCHEMA)
            rows = connection.execute(
                "SELECT portal_ref, portal_status, ean FROM marcolin_portal_groups WHERE request_id=? "
                "ORDER BY recorded_at, rowid", (request_id,)).fetchall()
        current = {(row["ean"], row["portal_ref"]): row["portal_status"] for row in rows}
        groups: dict[tuple, dict] = {}
        for (ean, ref), status in sorted(current.items(), key=lambda item: (item[0][1], item[0][0])):
            groups.setdefault((ref, status), {"portal_ref": ref, "status": status, "upcs": []})["upcs"].append(ean)
        return list(groups.values())

    def export_request_xlsx(self, request_id: int, *, engine=None, odoo_root=None) -> bytes:
        """Build the operator handoff workbook without changing request state.

        Odoo is consulted only when explicitly requested by the caller and only when this
        request has no persisted evidence yet. Missing or ambiguous matches remain visible.
        """
        request = self.service.store.get(request_id)
        items = self.service.store.odoo_items(request_id)
        if not items and engine is not None:
            from provider_odoo_orders import confirmed_partner_ids, read_confirmed_skus
            items = read_confirmed_skus(
                engine, confirmed_partner_ids(odoo_root, self.provider),
                [item["ean"] for item in request["eans"]])
            self.service.store.save_odoo_items(request_id, items)
        by_ean = {item["ean"]: item for item in items}
        portal_refs = {}
        for group in self.portal_groups(request_id):
            if len(group["upcs"]) == 1:
                portal_refs[group["upcs"][0]] = group["portal_ref"]
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Marcolin"
        headers = ["UPC / EAN", "SKU completo", "Código real", "Referencia fabricante confirmada",
                   "Producto Odoo", "Proveedor Odoo", "Pedido Odoo", "Fecha Odoo",
                   "Referencia portal", "Estado", "Motivo"]
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        for row in request["eans"]:
            item = by_ean.get(row["ean"], {})
            status = item.get("status", "pending")
            sheet.append([
                row["ean"], item.get("sku", "") if status == "confirmed" else "",
                row["ean"], item.get("sku", "") if status == "confirmed" else "",
                item.get("product", ""), item.get("supplier", ""), item.get("order_name", ""),
                item.get("order_date", ""), portal_refs.get(row["ean"], ""),
                status, item.get("reason", "sin correspondencia Odoo confirmada"),
            ])
        for cell in sheet["A"][1:]:
            cell.number_format = "@"
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        widths = [16, 24, 16, 30, 34, 24, 18, 14, 24, 14, 38]
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[chr(64 + index)].width = width
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()
